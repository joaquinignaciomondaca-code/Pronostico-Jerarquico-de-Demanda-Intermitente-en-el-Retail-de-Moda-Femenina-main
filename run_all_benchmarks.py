# -*- coding: utf-8 -*-
"""
Script Orquestador Maestro - Pipeline Completo de Tesis
Ejecuta secuencialmente todos los modelos base (Seasonal Naive, Croston, SBA, SARIMA, SARIMAX, LSTM, Chronos)
y la arquitectura propuesta ES-GBM con Learned Shares, consolidando todo en un ÚNICO archivo Excel maestro:
resultados/RESULTADOS_TESIS_CONSOLIDADO.xlsx
"""

import os
import sys
import time
import subprocess
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
RES_DIR = os.path.join(HERE, 'resultados')
MASTER_EXCEL = os.path.join(RES_DIR, 'RESULTADOS_TESIS_CONSOLIDADO.xlsx')
os.makedirs(RES_DIR, exist_ok=True)

SCRIPTS = [
    ("Seasonal Naive (Lag 12)", os.path.join(HERE, "01_modelos_base", "e)_seasonal_naive_topdown.py")),
    ("Croston Clásico (1972)", os.path.join(HERE, "01_modelos_base", "f)_croston_topdown.py")),
    ("SBA (Syntetos-Boylan 2005)", os.path.join(HERE, "01_modelos_base", "g)_sba_topdown.py")),
    ("SARIMA (1,1,1)(1,1,1)12", os.path.join(HERE, "01_modelos_base", "a)_sarima_topdown.py")),
    ("SARIMAX (con precio ex-ante)", os.path.join(HERE, "01_modelos_base", "b)_sarimax_topdown.py")),
    ("LSTM (PyTorch Multi-Seed)", os.path.join(HERE, "01_modelos_base", "c)_lstm_topdown.py")),
    ("Amazon Chronos-Bolt Small", os.path.join(HERE, "01_modelos_base", "d)_chronos_topdown.py")),
    ("ES-GBM Ensemble (Familia)", os.path.join(HERE, "02_modelo_final", "2)_esgbm_ensemble.py")),
    ("Learned Shares (SKU)", os.path.join(HERE, "02_modelo_final", "3)_desagregacion_sku_share.py"))
]

def format_excel_workbook(filepath):
    wb = openpyxl.load_workbook(filepath)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    regular_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='E0E0E0'),
        right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='E0E0E0'),
        bottom=Side(style='thin', color='E0E0E0')
    )

    for ws in wb.worksheets:
        ws.views.sheetView[0].showGridLines = True
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = regular_font
                cell.border = thin_border
                if isinstance(cell.value, float):
                    col_name = str(ws.cell(1, cell.column).value).lower()
                    if any(k in col_name for k in ['wape', '%', 'bias', 'share']):
                        cell.number_format = '0.00%'
                    else:
                        cell.number_format = '#,##0.00'
                elif isinstance(cell.value, int) and cell.value > 1000:
                    cell.number_format = '#,##0'
                    
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 13), 45)

    wb.save(filepath)
    wb.close()

def main():
    print("=" * 80)
    print("   EJECUCIÓN INTEGRAL DEL PIPELINE DE MODELOS PREDICTIVOS - UTFSM 2026")
    print("=" * 80)
    start_total = time.time()
    
    for idx, (name, script_path) in enumerate(SCRIPTS, 1):
        print(f"\n[{idx}/{len(SCRIPTS)}] Ejecutando: {name}...")
        print("-" * 80)
        t0 = time.time()
        
        result = subprocess.run([sys.executable, script_path], cwd=HERE)
        elapsed = time.time() - t0
        
        if result.returncode == 0:
            print(f" -> [OK] Finalizado exitosamente en {elapsed:.1f}s")
        else:
            print(f" -> [ERROR] Falló la ejecución de {name} (código {result.returncode})")
            
    print("\n" + "=" * 80)
    print("   FORMATEANDO LIBRO MAESTRO DE RESULTADOS CONSOLIDADO")
    print("=" * 80)
    
    if os.path.exists(MASTER_EXCEL):
        format_excel_workbook(MASTER_EXCEL)
        print(f" -> [OK] Archivo maestro actualizado y estilizado en: {MASTER_EXCEL}")
        
        df_summary = pd.read_excel(MASTER_EXCEL, sheet_name='00_Resumen_Consolidado')
        print("\n=== TABLA OFICIAL CONSOLIDADA ===")
        print(df_summary.to_string(index=False))
        
    print("-" * 80)
    print(f"Tiempo total de ejecución del pipeline: {(time.time() - start_total):.1f}s")
    print("=" * 80)

if __name__ == '__main__':
    main()
